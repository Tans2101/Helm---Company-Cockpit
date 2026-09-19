"""Income Statement + Cash Summary financial export."""
import os
import sys
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path
from unittest.mock import AsyncMock, MagicMock, patch

import pytest
from fastapi.testclient import TestClient

os.environ.setdefault("MONGO_URL", "mongodb://localhost:27017")
os.environ.setdefault("DB_NAME", "test_financial_export")

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

import financial_export as fin_exp  # noqa: E402
import finance_recurrence as fin_recur  # noqa: E402
import server  # noqa: E402

NOW = datetime(2026, 9, 8, tzinfo=timezone.utc)

ENTRIES = [
    {
        "id": "r1",
        "type": "revenue",
        "category": "Subscriptions",
        "name": "Subscription MRR",
        "amount": 10000,
        "month": "2026-08",
        "recurring": True,
    },
    {
        "id": "e1",
        "type": "expense",
        "category": "Payroll",
        "name": "Team payroll",
        "amount": 4000,
        "month": "2026-08",
        "recurring": True,
    },
    {
        "id": "e2",
        "type": "expense",
        "category": "Cloud/Infra",
        "name": "MongoDB Database Subscription",
        "amount": 1000,
        "month": "2026-08",
        "recurring": False,
    },
]
SETTINGS = {"cash": 50000, "cash_entered": True, "currency": "usd"}


def test_assemble_matches_dashboard_expansion():
    bundle = fin_exp.assemble_financial_export(ENTRIES, SETTINGS, "2026-09", now=NOW)
    horizon = fin_recur.resolve_expense_horizon(ENTRIES, NOW)
    rev = fin_recur.expand_entries_by_month(ENTRIES, entry_type="revenue", horizon_end=horizon)
    exp = fin_recur.expand_entries_by_month(ENTRIES, entry_type="expense", horizon_end=horizon)
    assert bundle["income"]["revenue"] == pytest.approx(rev["2026-09"])
    assert bundle["income"]["expenses_total"] == pytest.approx(exp["2026-09"])
    assert bundle["income"]["net_income"] == pytest.approx(rev["2026-09"] - exp["2026-09"])
    # Dashboard cash snapshot is ending cash of the latest month
    assert bundle["latest_month"] == "2026-09"
    assert bundle["cash"]["ending"] == pytest.approx(50000)
    assert bundle["cash"]["ending_matches_dashboard"] is True
    assert bundle["cash"]["inflows"] == pytest.approx(rev["2026-09"])
    assert bundle["cash"]["outflows"] == pytest.approx(exp["2026-09"])
    # Starting = ending - inflows + outflows
    assert bundle["cash"]["starting"] == pytest.approx(50000 - rev["2026-09"] + exp["2026-09"])


def test_august_walks_back_from_dashboard_cash():
    sept = fin_exp.assemble_financial_export(ENTRIES, SETTINGS, "2026-09", now=NOW)
    aug = fin_exp.assemble_financial_export(ENTRIES, SETTINGS, "2026-08", now=NOW)
    assert aug["cash"]["ending"] == pytest.approx(sept["cash"]["starting"])
    assert aug["income"]["revenue"] == pytest.approx(10000)
    assert aug["income"]["expenses_total"] == pytest.approx(5000)
    cats = {row["category"]: row["amount"] for row in aug["income"]["expenses_by_category"]}
    assert cats["Payroll"] == pytest.approx(4000)
    assert cats["Cloud/Infra"] == pytest.approx(1000)
    names = {row["name"]: row for row in aug["line_items"]}
    assert names["Team payroll"]["category"] == "Payroll"
    assert names["MongoDB Database Subscription"]["category"] == "Cloud/Infra"
    md = fin_exp.statement_markdown(aug)
    assert "MongoDB Database Subscription" in md
    assert "(Cloud/Infra, expense)" in md


def test_missing_cash_is_not_zero():
    bundle = fin_exp.assemble_financial_export(ENTRIES, {"currency": "usd"}, "2026-09", now=NOW)
    assert bundle["cash"]["entered"] is False
    assert bundle["cash"]["starting"] is None
    assert bundle["cash"]["ending"] is None
    assert bundle["cash"]["dashboard_cash"] is None


def test_no_balance_sheet_language_in_exports():
    bundle = fin_exp.assemble_financial_export(ENTRIES, SETTINGS, "2026-09", now=NOW)
    md = fin_exp.statement_markdown(bundle)
    assert "balance sheet" in md.lower()
    assert "no balance sheet" in md.lower()
    src = Path(fin_exp.__file__).read_text()
    assert "Balance Sheet" not in src
    xlsx = fin_exp.render_financial_xlsx(bundle, workspace_name="Acme")
    from openpyxl import load_workbook

    wb = load_workbook(BytesIO(xlsx))
    assert wb.sheetnames == ["Income Statement", "Cash Summary", "Line items"]
    assert "Balance" not in "".join(wb.sheetnames)


def test_pdf_reuses_weekly_pack_pipeline():
    bundle = fin_exp.assemble_financial_export(ENTRIES, SETTINGS, "2026-09", now=NOW)
    pdf = fin_exp.render_financial_pdf(bundle, workspace_name="Northwind")
    assert pdf.startswith(b"%PDF")
    assert b"%%EOF" in pdf[-64:]
    src = Path(fin_exp.__file__).read_text()
    assert "render_document_pdf" in src
    assert "SimpleDocTemplate" not in src


def test_xlsx_is_formatted_workbook_not_csv():
    bundle = fin_exp.assemble_financial_export(ENTRIES, SETTINGS, "2026-09", now=NOW)
    raw = fin_exp.render_financial_xlsx(bundle, workspace_name="Forge Co")
    assert raw[:2] == b"PK"
    from openpyxl import load_workbook

    wb = load_workbook(BytesIO(raw))
    income = wb["Income Statement"]
    cash = wb["Cash Summary"]
    assert income["A5"].value == "Line"
    assert income["B5"].value == "Amount"
    # Revenue row is a real number with accounting format
    assert income["B6"].value == pytest.approx(bundle["income"]["revenue"])
    assert income["B6"].number_format == "#,##0.00"
    assert cash["B7"].value == pytest.approx(bundle["cash"]["inflows"])
    assert cash["B7"].number_format == "#,##0.00"


def test_invalid_period_rejected():
    with pytest.raises(ValueError):
        fin_exp.assemble_financial_export(ENTRIES, SETTINGS, "2026-13", now=NOW)


def _owner_principal():
    async def mock_principal():
        return {
            "user_id": "u_owner",
            "email": "ceo@example.com",
            "name": "CEO",
            "workspace_id": "ws_1",
            "role": "owner",
            "pack": "owner",
        }

    return mock_principal


def _patch_export_deps():
    ws = {
        "workspace_id": "ws_1",
        "name": "Forge Co",
        "plan": "starter",
        "financial_settings": SETTINGS,
    }

    class _Find:
        def __init__(self, rows):
            self._rows = rows

        async def to_list(self, _n):
            return list(self._rows)

    mock_db = MagicMock()
    mock_db.financial_entries.find.return_value = _Find(ENTRIES)

    return (
        patch.object(server, "get_ws", new=AsyncMock(return_value=ws)),
        patch.object(server, "BILLING_ENFORCED", False),
        patch.object(server, "db", mock_db),
        patch.object(server.dept_access, "accessible_department_ids", new=AsyncMock(return_value=None)),
        patch.object(server.dept_access, "apply_department_filter", side_effect=lambda filt, _ids: filt),
    )


def test_pdf_endpoint_returns_attachment():
    server.app.dependency_overrides[server.get_principal] = _owner_principal()
    patches = _patch_export_deps()
    try:
        with patches[0], patches[1], patches[2], patches[3], patches[4]:
            client = TestClient(server.app)
            r = client.post("/api/reports/financial-export/pdf", json={"period": "2026-09"})
    finally:
        server.app.dependency_overrides.clear()
    assert r.status_code == 200, r.text
    assert r.headers["content-type"].startswith("application/pdf")
    assert "Trenston-Financial-Export-Forge-Co-2026-09.pdf" in r.headers.get("content-disposition", "")
    assert r.content.startswith(b"%PDF")


def test_xlsx_endpoint_returns_attachment():
    server.app.dependency_overrides[server.get_principal] = _owner_principal()
    patches = _patch_export_deps()
    try:
        with patches[0], patches[1], patches[2], patches[3], patches[4]:
            client = TestClient(server.app)
            r = client.post("/api/reports/financial-export/xlsx", json={"period": "2026-08"})
    finally:
        server.app.dependency_overrides.clear()
    assert r.status_code == 200, r.text
    assert "spreadsheetml" in r.headers["content-type"]
    assert r.content[:2] == b"PK"
    from openpyxl import load_workbook

    wb = load_workbook(BytesIO(r.content))
    assert set(wb.sheetnames) == {"Income Statement", "Cash Summary", "Line items"}
    li = wb["Line items"]
    names = [cell.value for cell in li["A"] if cell.value]
    assert "MongoDB Database Subscription" in names
    assert "Cloud/Infra" in [cell.value for cell in li["B"] if cell.value]


def test_export_forbidden_without_finance_write():
    async def mock_principal():
        return {
            "user_id": "u_mem",
            "email": "member@example.com",
            "workspace_id": "ws_1",
            "role": "member",
            "pack": "member",
        }

    server.app.dependency_overrides[server.get_principal] = mock_principal
    try:
        with patch.object(server, "BILLING_ENFORCED", False), \
             patch.object(server, "_membership_for", new=AsyncMock(return_value={})), \
             patch.object(server, "get_ws", new=AsyncMock(return_value={"workspace_id": "ws_1", "section_access": {}})):
            client = TestClient(server.app)
            pdf = client.post("/api/reports/financial-export/pdf", json={"period": "2026-09"})
            xlsx = client.post("/api/reports/financial-export/xlsx", json={"period": "2026-09"})
    finally:
        server.app.dependency_overrides.clear()
    assert pdf.status_code == 403
    assert xlsx.status_code == 403


def test_export_rejects_bad_period():
    server.app.dependency_overrides[server.get_principal] = _owner_principal()
    patches = _patch_export_deps()
    try:
        with patches[0], patches[1], patches[2], patches[3], patches[4]:
            client = TestClient(server.app)
            r = client.post("/api/reports/financial-export/pdf", json={"period": "not-a-month"})
    finally:
        server.app.dependency_overrides.clear()
    assert r.status_code == 400
