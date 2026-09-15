"""Convert uploaded report spreadsheets into plain text for LLM summarization."""
from __future__ import annotations

import csv
import io
from typing import Tuple

REPORT_TEXT_CHAR_BUDGET = 80_000
REPORT_MAX_ROWS = 400

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
CSV_MIMES = frozenset({"text/csv", "application/csv"})
SPREADSHEET_MIMES = frozenset({XLSX_MIME, *CSV_MIMES})


def is_spreadsheet_mime(content_type: str) -> bool:
    return (content_type or "") in SPREADSHEET_MIMES


def spreadsheet_bytes_to_text(data: bytes, content_type: str, filename: str = "") -> Tuple[str, bool]:
    """Return (markdown-ish table text, truncated).

    Caps rows and characters so a large sheet cannot blow the model context.
    """
    ctype = content_type or ""
    if ctype in CSV_MIMES or (filename or "").lower().endswith(".csv"):
        return _csv_to_text(data)
    if ctype == XLSX_MIME or (filename or "").lower().endswith(".xlsx"):
        return _xlsx_to_text(data)
    raise ValueError(f"Unsupported spreadsheet type: {ctype or filename or 'unknown'}")


def _csv_to_text(data: bytes) -> Tuple[str, bool]:
    for encoding in ("utf-8-sig", "utf-8", "latin-1"):
        try:
            text = data.decode(encoding)
            break
        except UnicodeDecodeError:
            text = None
    if text is None:
        text = data.decode("utf-8", errors="replace")

    reader = csv.reader(io.StringIO(text))
    lines: list[str] = []
    truncated = False
    for i, row in enumerate(reader):
        if i >= REPORT_MAX_ROWS:
            truncated = True
            break
        cells = [str(c).replace("\n", " ").strip() for c in row]
        lines.append(" | ".join(cells))
        if sum(len(x) + 1 for x in lines) >= REPORT_TEXT_CHAR_BUDGET:
            truncated = True
            break
    body = "\n".join(lines).strip()
    if len(body) > REPORT_TEXT_CHAR_BUDGET:
        body = body[:REPORT_TEXT_CHAR_BUDGET]
        truncated = True
    if truncated:
        body = (body.rstrip() + "\n\n[Truncated: only the first rows of this CSV were included.]").strip()
    return body or "(empty CSV)", truncated


def _xlsx_to_text(data: bytes) -> Tuple[str, bool]:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    parts: list[str] = []
    truncated = False
    rows_used = 0
    chars_used = 0

    for sheet in wb.worksheets:
        parts.append(f"## Sheet: {sheet.title}")
        chars_used += len(parts[-1]) + 1
        for row in sheet.iter_rows(values_only=True):
            if rows_used >= REPORT_MAX_ROWS or chars_used >= REPORT_TEXT_CHAR_BUDGET:
                truncated = True
                break
            cells = []
            empty = True
            for cell in row:
                if cell is None:
                    cells.append("")
                else:
                    empty = False
                    cells.append(str(cell).replace("\n", " ").strip())
            if empty:
                continue
            line = " | ".join(cells)
            parts.append(line)
            rows_used += 1
            chars_used += len(line) + 1
        if truncated:
            break

    try:
        wb.close()
    except Exception:
        pass

    body = "\n".join(parts).strip()
    if len(body) > REPORT_TEXT_CHAR_BUDGET:
        body = body[:REPORT_TEXT_CHAR_BUDGET]
        truncated = True
    if truncated:
        body = (body.rstrip() + "\n\n[Truncated: only the first rows of this spreadsheet were included.]").strip()
    return body or "(empty spreadsheet)", truncated
